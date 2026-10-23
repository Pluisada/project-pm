"""Part A: excel_contract + commercial_excel_reader. Workbooks are synthetic and
built in memory; no database, no real data."""
import io
import zipfile
from datetime import date, datetime, time, timedelta
from decimal import Decimal

import pytest
from openpyxl import Workbook

import excel_contract as contract
from bid_excel_reader import ReaderLimits
from commercial_excel_reader import (
    FORMULA,
    _FieldError,
    parse_money,
    read_commercial_excel,
)

HEADERS = [spec.header for spec in contract.COLUMNS]  # every accepted column, legacy ones included


def make(rows, headers=None, sheet="Cards", extra_sheets=(), merge=None):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers if headers is not None else HEADERS)
    for r in rows:
        ws.append(r)
    if merge:
        ws.merge_cells(merge)
    for name in extra_sheets:
        wb.create_sheet(name)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def good(**overrides):
    values = {"title": "Lead A", "company": "Acme", "deal_value_usd": 1000}
    values.update(overrides)
    return [values.get(h) for h in HEADERS]


def read(data, name="cards.xlsx", **kw):
    return read_commercial_excel(data, name, **kw)


def codes(result):
    return [e.code for r in result.rows for e in r.errors]


def one(result):
    assert result.ok, result.file_errors
    return result.rows[0]


# --- contract ---------------------------------------------------------------

def test_contract_template_headers_are_exactly_the_expected_ones_without_forbidden_ones():
    assert contract.HEADERS == (
        "card_id", "title", "company", "contact", "contact_email", "contact_phone", "details", "deal_value_usd",
        "deal_value_next_year_usd", "next_action_date", "probability_level", "product_line", "lead_source",
        "salesperson_email", "stage",
    )
    # Derived or obsolete inputs are not in the template, but older workbooks carrying them are still read.
    for legacy in ("deal_total_usd", "deal_year", "expected_close_date"):
        assert legacy not in contract.HEADERS and legacy in contract.SPECS_BY_KEY
    assert not any(contract.is_forbidden_header(h) for h in HEADERS)
    assert contract.is_forbidden_header("next_step_due")  # the Next Action Date column is next_action_date
    assert len(set(HEADERS)) == len(HEADERS)


def test_template_workbook_without_legacy_columns_reads_cleanly():
    result = read(make([[None, "Lead"] + [None] * (len(contract.HEADERS) - 2)], headers=list(contract.HEADERS)))
    assert result.file_errors == () and not any(w.column in ("deal_total_usd", "deal_year", "expected_close_date")
                                                for w in result.file_warnings)


def test_contact_fields_are_separate_and_validated_like_the_card_screen():
    row = one(read(make([good(contact="Ana", contact_email=" ana@acme.com ", contact_phone="+55 (11) 99999-0000")])))
    rec = row.record
    assert (rec.contact, rec.contact_email, rec.contact_phone) == ("Ana", "ana@acme.com", "+55 (11) 99999-0000")
    assert codes(read(make([good(contact_email="nope")]))) == ["invalid_email"]
    assert codes(read(make([good(contact_phone="call me")]))) == ["invalid_phone"]


def test_next_action_date_is_read_into_the_next_step_due_field():
    assert one(read(make([good(next_action_date="2026-12-31")]))).record.next_step_due == date(2026, 12, 31)
    assert one(read(make([good(next_action_date=datetime(2026, 5, 1))]))).record.next_step_due == date(2026, 5, 1)
    for bad in ("2026-02-30", "31/12/2026", "tomorrow"):
        assert codes(read(make([good(next_action_date=bad)]))) == ["invalid_date"], bad
    assert one(read(make([good(expected_close_date="2026-12-31")]))).record.next_step_due is None


# --- headers ----------------------------------------------------------------

def test_valid_file_and_free_column_order():
    reordered = list(reversed(HEADERS))
    values = {"title": "T", "stage": "RFI"}
    result = read(make([[values.get(h) for h in reordered]], headers=[h.upper() + "  " for h in reordered]))
    row = one(result)
    assert row.is_valid and row.record.title == "T" and row.record.stage == "RFI"
    assert result.file_warnings == ()


def test_missing_title_column_is_file_error():
    headers = [h for h in HEADERS if h != "title"]
    result = read(make([[None] * len(headers)], headers=headers))
    assert [e.code for e in result.file_errors] == ["missing_column"] and result.rows == ()


def test_missing_optional_column_is_a_warning():
    headers = ["card_id", "title", "company"]
    result = read(make([[None, "T", "C"]], headers=headers))
    assert result.ok and one(result).is_valid
    assert {w.code for w in result.file_warnings} == {"missing_optional_column"}


def test_card_id_column_is_required_even_when_every_cell_would_be_blank():
    result = read(make([["T"]], headers=["title"]))
    assert [(e.code, e.column) for e in result.file_errors] == [("missing_column", "card_id")] and result.rows == ()


def test_duplicate_header_is_file_error():
    result = read(make([["a", "b"]], headers=["title", " TITLE"]))
    assert "duplicate_header" in [e.code for e in result.file_errors]


@pytest.mark.parametrize("header", ["next_step", "next_step_due", "revenue_current_year", "Bid_Code"])
def test_forbidden_columns_are_file_errors(header):
    result = read(make([[None, "T", "x"]], headers=["card_id", "title", header]))
    assert [e.code for e in result.file_errors] == ["forbidden_column"]
    assert result.file_errors[0].column == header


def test_unknown_column_is_file_error_not_silently_ignored():
    result = read(make([[None, "T", "x"]], headers=["card_id", "title", "budget"]))
    assert [e.code for e in result.file_errors] == ["unknown_column"]


def test_blank_header_in_the_middle_is_file_error():
    result = read(make([["T", "x", "y"]], headers=["title", None, "company"]))
    assert "blank_header" in [e.code for e in result.file_errors]


# --- sheets -----------------------------------------------------------------

def test_missing_cards_sheet():
    result = read(make([good()], sheet="Other"))
    assert [e.code for e in result.file_errors] == ["sheet_not_found"]


def test_sheet_name_matches_ignoring_case_and_extra_sheets_warn():
    result = read(make([good()], sheet="cards", extra_sheets=("Notes",)))
    assert result.ok and [w.code for w in result.file_warnings] == ["unexpected_sheet"]
    assert "Notes" not in result.file_warnings[0].message


def test_merged_cells_rejected():
    result = read(make([good()], merge="C2:D2"))
    assert [e.code for e in result.file_errors] == ["merged_cells"]


def test_no_data_rows():
    assert [e.code for e in read(make([])).file_errors] == ["no_data_rows"]


def test_blank_rows_are_skipped():
    result = read(make([good(), [None] * 15, ["  "] + [None] * 14, good(title="B")]))
    assert len(result.rows) == 2 and result.blank_rows_skipped == 2
    assert [r.row_number for r in result.rows] == [2, 5]


# --- file safety --------------------------------------------------------------

def test_malformed_and_wrong_files():
    assert read(b"not a zip at all", "x.xlsx").file_errors[0].code == "not_xlsx"
    assert read(b"PK\x03\x04garbage", "x.xlsx").file_errors[0].code == "corrupted_file"
    assert read(b"", "x.xlsx").file_errors[0].code == "empty_file"
    assert read(make([good()]), "x.csv").file_errors[0].code == "invalid_extension"
    assert read(make([good()]), "x.xlsm").file_errors[0].code == "format_not_allowed"


def test_real_content_not_just_extension():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "hi")
    assert read(buf.getvalue(), "x.xlsx").file_errors[0].code == "not_xlsx"


def test_macro_part_rejected():
    src = make([good()])
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(src)) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("xl/vbaProject.bin", b"x")
    assert read(out.getvalue()).file_errors[0].code == "format_not_allowed"


def test_zip_bomb_rejected():
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(make([good()]))) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item.filename, zin.read(item.filename))
        zout.writestr("xl/bomb.bin", b"\0" * (30 * 1024 * 1024))
    assert read(out.getvalue()).file_errors[0].code == "suspicious_compression"


def test_xml_entities_rejected():
    src = make([good()])
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(src)) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "b">]>' + data
            zout.writestr(item, data)
    assert read(out.getvalue()).file_errors[0].code == "unsafe_xml"


def test_size_and_row_limits():
    assert read(make([good()]), limits=ReaderLimits(max_file_bytes=100)).file_errors[0].code == "file_too_large"
    result = read(make([good(title=f"T{i}") for i in range(4)]), limits=ReaderLimits(max_rows=3))
    assert [e.code for e in result.file_errors] == ["too_many_rows"] and result.rows == ()
    assert read(make([good(title=f"T{i}") for i in range(3)]), limits=ReaderLimits(max_rows=3)).ok
    assert contract and ReaderLimits().max_rows == 5000 and ReaderLimits().max_file_bytes == 10 * 1024 * 1024


def test_formulas_are_rejected_even_with_cached_value_or_without():
    data = make([good(title="=1+1"), good(deal_value_usd="=SUM(1,2)")])
    result = read(data)
    assert codes(result) == ["formula_not_allowed", "formula_not_allowed"]
    assert result.rows[0].record.title is None


def test_formula_with_cached_value_is_rejected():
    # a cached value present in the XML must not make the formula acceptable
    src = make([good()])
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(src)) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'<v>1000</v>', b'<f>500+500</f><v>1000</v>', 1)
            zout.writestr(item, data)
    assert codes(read(out.getvalue())) == ["formula_not_allowed"]


# --- fields -------------------------------------------------------------------

def test_card_id_blank_is_a_new_card_and_filled_is_an_update_reference():
    result = read(make([good(card_id=7), good(title="B"), good(card_id=" 12 ", title="C")]))
    assert codes(result) == []
    assert [row.record.card_id for row in result.rows] == [7, None, 12]
    assert "card_id" in result.columns and "title" in result.columns


@pytest.mark.parametrize("value", [0, -3, 1.5, "abc", "1.0e3", "-1", "٣", 2**64])
def test_card_id_must_be_a_positive_whole_number(value):
    result = read(make([good(card_id=value), good(title="B")]))
    assert codes(result) == ["invalid_card_id"]
    assert result.rows[0].record.card_id is None and result.rows[1].is_valid


def test_card_id_formula_is_rejected():
    assert codes(read(make([good(card_id="=1+1")]))) == ["formula_not_allowed"]


def test_title_required_trimmed_and_limited():
    assert codes(read(make([good(title=None)]))) == ["required_value_missing"]
    assert codes(read(make([good(title="   ")]))) == ["required_value_missing"]
    assert one(read(make([good(title="  X  ")]))).record.title == "X"
    assert one(read(make([good(title="x" * 255)]))).is_valid
    assert codes(read(make([good(title="x" * 256)]))) == ["text_too_long"]


@pytest.mark.parametrize("field,limit", [("company", 255), ("contact", 255), ("details", 5000)])
def test_text_limits(field, limit):
    assert one(read(make([good(**{field: "y" * limit})]))).is_valid
    assert codes(read(make([good(**{field: "y" * (limit + 1)})]))) == ["text_too_long"]


def test_error_messages_never_contain_cell_content():
    secret = "SEGREDO-" + "z" * 300
    result = read(make([good(company=secret, contact="x@y", lead_source=secret, probability_level=secret,
                              deal_value_usd=secret, expected_close_date=secret, salesperson_email=secret)]))
    assert result.rows[0].errors
    for e in result.rows[0].errors:
        assert "SEGREDO" not in e.message and e.column and e.row == 2


def test_money_empty_is_null_zero_is_zero_and_next_year_defaults_to_zero():
    row = one(read(make([good(deal_value_usd=None)]))).record
    assert row.deal_value_usd is None and row.deal_value_next_year_usd == 0
    row = one(read(make([good(deal_value_usd=0, deal_value_next_year_usd=0)]))).record
    assert row.deal_value_usd == Decimal("0.00") and row.deal_value_usd is not None


def test_money_decimal_exact_and_two_places():
    row = one(read(make([good(deal_value_usd=Decimal("1234.56"), deal_value_next_year_usd="10.5")]))).record
    assert row.deal_value_usd == Decimal("1234.56") and row.deal_value_next_year_usd == Decimal("10.50")


@pytest.mark.parametrize("value,code", [
    (-1, "negative_amount"), ("-5", "negative_amount"), (1.234, "too_many_decimals"),
    ("US$ 1.000", "invalid_amount"), ("1,000.00", "invalid_amount"), ("abc", "invalid_amount"),
    (10 ** 13, "amount_too_large"), (True, "invalid_type"),
])
def test_money_rejections(value, code):
    assert codes(read(make([good(deal_value_usd=value)]))) == [code]


@pytest.mark.parametrize("value", [Decimal("NaN"), Decimal("Infinity"), Decimal("-Infinity"), Decimal("sNaN")])
def test_money_non_finite(value):
    with pytest.raises(_FieldError) as exc:
        parse_money(value)
    assert exc.value.code == "invalid_amount"


def test_money_limit_boundary():
    assert parse_money(contract.MAX_DEAL_VALUE) == contract.MAX_DEAL_VALUE
    with pytest.raises(_FieldError):
        parse_money(contract.MAX_DEAL_VALUE + Decimal("0.01"))


def test_total_is_ignored_and_never_stored():
    ok = one(read(make([good(deal_value_usd=100, deal_value_next_year_usd=50, deal_total_usd=150)])))
    assert ok.is_valid and not hasattr(ok.record, "deal_total_usd")


def test_stale_total_does_not_block_and_annual_values_are_kept():
    row = one(read(make([good(deal_value_usd=60000, deal_value_next_year_usd=20000, deal_total_usd=75000)])))
    assert row.is_valid
    assert row.record.deal_value_usd == Decimal("60000.00")
    assert row.record.deal_value_next_year_usd == Decimal("20000.00")


def test_total_filled_with_empty_annual_values():
    row = one(read(make([good(deal_value_usd=None, deal_value_next_year_usd=None, deal_total_usd=500)])))
    assert row.is_valid
    assert row.record.deal_value_usd is None and row.record.deal_value_next_year_usd == 0


def test_total_filled_without_annual_columns():
    headers = ["card_id", "title", "deal_total_usd"]
    result = read(make([[None, "Lead A", 500]], headers=headers))
    row = one(result)
    assert row.is_valid and row.record.deal_value_usd is None and row.record.deal_value_next_year_usd == 0


@pytest.mark.parametrize("bad_total", ["abc", "-5", "--5", 1.234, 10 ** 13, True, time(10, 30), timedelta(hours=2)])
def test_invalid_total_content_does_not_block(bad_total):
    row = one(read(make([good(deal_value_usd=100, deal_value_next_year_usd=50, deal_total_usd=bad_total)])))
    assert row.is_valid and row.record.deal_value_usd == 100


@pytest.mark.parametrize("value", ["--5", "---5", "-", "+5"])
def test_malformed_signed_text_is_a_row_error_not_corrupted_file(value):
    result = read(make([good(deal_value_usd=value)]))
    assert result.ok and codes(result) == ["invalid_amount"]


@pytest.mark.parametrize("value", [time(10, 30), timedelta(hours=2)])
def test_time_or_duration_in_money_column_is_a_row_error(value):
    result = read(make([good(deal_value_usd=value)]))
    assert result.ok and codes(result) == ["invalid_type"]
    assert result.rows[0].errors[0].column == "deal_value_usd" and result.rows[0].errors[0].row == 2


@pytest.mark.parametrize("value", [time(10, 30), timedelta(hours=2)])
def test_time_or_duration_in_text_column_is_a_row_error(value):
    result = read(make([good(company=value)]))
    assert result.ok and codes(result) == ["invalid_type"]
    assert result.rows[0].errors[0].column == "company" and result.rows[0].errors[0].row == 2


def test_deal_year_is_informational_only():
    row = one(read(make([good(deal_year=2027)])))
    assert row.record.deal_year_hint == 2027 and not hasattr(row.record, "deal_year")
    assert codes(read(make([good(deal_year=1800)]))) == ["invalid_year"]
    assert codes(read(make([good(deal_year="soon")]))) == ["invalid_year"]


def test_probability_level():
    for text in ("low", "Medium", " HIGH "):
        assert one(read(make([good(probability_level=text)]))).record.confidence_level == text.strip().upper()
    assert one(read(make([good(probability_level=None)]))).record.confidence_level is None
    assert codes(read(make([good(probability_level="URGENT")]))) == ["invalid_probability_level"]


def test_dates():
    assert one(read(make([good(expected_close_date="2026-12-31")]))).record.expected_close_date == date(2026, 12, 31)
    assert one(read(make([good(expected_close_date=datetime(2026, 5, 1))]))).record.expected_close_date == date(2026, 5, 1)
    assert one(read(make([good(expected_close_date=date(2026, 5, 1))]))).record.expected_close_date == date(2026, 5, 1)
    for bad in ("2026-02-30", "31/12/2026", "12/31/2026", "2026/12/31", "tomorrow", 46000, "2026-1-5"):
        assert codes(read(make([good(expected_close_date=bad)]))) == ["invalid_date"], bad


def test_lead_source_uses_existing_values_case_insensitively():
    from lead_options import LEAD_SOURCES

    for source in LEAD_SOURCES:
        assert one(read(make([good(lead_source=source.upper())]))).record.lead_source == source
    assert codes(read(make([good(lead_source="Billboard")]))) == ["invalid_lead_source"]


def test_syntax_only_context_fields_are_passed_through():
    row = one(read(make([good(product_line="Anything", salesperson_email="  Ana@Example.COM ", stage="Whatever")])))
    assert row.is_valid
    assert (row.record.product_line, row.record.salesperson_email, row.record.stage) == (
        "Anything", "ana@example.com", "Whatever")
    assert codes(read(make([good(salesperson_email="not-an-email")]))) == ["invalid_email"]
    assert all(contract.SPECS_BY_KEY[k].context for k in ("product_line", "salesperson_email", "stage"))


def test_value_outside_headers_is_row_error():
    data = make([good() + ["stray"]])
    assert codes(read(data)) == ["value_outside_headers"]


def test_each_error_has_row_column_code_and_scope():
    result = read(make([good(), good(title=None)]))
    e = result.rows[1].errors[0]
    assert (e.scope, e.code, e.row, e.column) == ("row", "required_value_missing", 3, "title")
    assert result.valid_rows == 1 and result.invalid_rows == 1


def test_reader_is_deterministic_and_has_no_side_effects(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    data = make([good()])
    assert read(data) == read(data)
    assert list(tmp_path.iterdir()) == []


def test_formula_marker_is_not_a_cached_value():
    assert FORMULA is not None
