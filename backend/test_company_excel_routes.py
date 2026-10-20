"""Excel export and import template routes: Super Admin or the company's own Admin, role checked before any company database opens.

Disposable tmp_path identity and company databases only.
"""
import io

import pytest
from openpyxl import load_workbook

from test_company_admin_products import company_admin
from test_company_card_writes import admin, member
from test_company_card_writes import rig  # noqa: F401  (fixture)

ROUTES = ["/excel/export", "/excel/template"]


@pytest.mark.parametrize("path", ROUTES)
def test_super_admin_downloads_workbook(rig, path):
    r = rig.call("get", path, admin(rig))
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert r.headers["cache-control"] == "no-store"
    assert load_workbook(io.BytesIO(r.content)).sheetnames


def test_export_contains_cards_and_template_is_empty(rig):
    def data_rows(path):
        ws = load_workbook(io.BytesIO(rig.call("get", path, admin(rig)).content)).worksheets[0]
        return [row for row in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in row)]

    assert len(data_rows("/excel/export")) == 2
    assert data_rows("/excel/template") == []


@pytest.mark.parametrize("path", ROUTES)
@pytest.mark.parametrize("role", ["sales_member", "visitor"])
def test_other_roles_get_403_before_company_db_opens(rig, path, role):
    _, token = member(rig, role=role)
    before = rig.digest("lity")
    assert rig.call("get", path, token).status_code == 403
    assert rig.opened == [] and rig.digest("lity") == before


@pytest.mark.parametrize("path", ROUTES)
def test_unauthenticated_is_401(rig, path):
    assert rig.call("get", path).status_code == 401
    assert rig.opened == []


@pytest.mark.parametrize("path", ROUTES)
def test_user_without_membership_cannot_use_another_company(rig, path):
    _, token = member(rig, "m@example.com", role="sales_member", slug="lity")
    before = {s: rig.digest(s) for s in ("salcomp", "valor-log")}
    for other in ("salcomp", "valor-log"):
        assert rig.call("get", path, token, slug=other).status_code == 403
    assert rig.opened == [] and {s: rig.digest(s) for s in before} == before


def test_filename_is_server_built_and_ignores_filename_param(rig):
    r = rig.call("get", "/excel/export", admin(rig), params={"filename": "../evil.xlsx"})
    assert r.status_code == 200
    disposition = r.headers["content-disposition"]
    assert "evil" not in disposition and "cards-lity-board-1-" in disposition
    t = rig.call("get", "/excel/template", admin(rig), params={"filename": "../evil.xlsx"})
    assert t.headers["content-disposition"] == 'attachment; filename="cards-import-template.xlsx"'


def test_export_of_unknown_board_is_404(rig):
    r = rig.client.get("/api/companies/lity/boards/999/excel/export",
                       headers={"Authorization": f"Bearer {admin(rig)}"})
    assert r.status_code == 404


# --- company Admin (Part 2D-G12): own company only ---------------------------------------------------

@pytest.mark.parametrize("path", ROUTES)
def test_company_admin_downloads_own_company_workbook_and_template(rig, path):
    _, token = company_admin(rig)
    r = rig.call("get", path, token)
    assert r.status_code == 200 and r.headers["cache-control"] == "no-store"
    assert load_workbook(io.BytesIO(r.content)).sheetnames


def test_company_admin_export_keeps_card_ids_for_update_imports(rig):
    _, token = company_admin(rig)
    ws = load_workbook(io.BytesIO(rig.call("get", "/excel/export", token).content)).worksheets[0]
    header = [c.value for c in ws[1]]
    ids = [row[header.index("card_id")] for row in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in row)]
    assert header[0] == "card_id" and sorted(ids) == [1, 2]


@pytest.mark.parametrize("path", ROUTES)
def test_company_admin_cannot_use_another_company(rig, path):
    _, token = company_admin(rig, slug="lity")
    before = {s: rig.digest(s) for s in ("salcomp", "valor-log")}
    rig.opened.clear()
    for other in ("salcomp", "valor-log"):
        assert rig.call("get", path, token, slug=other).status_code == 403
    assert rig.opened == [] and {s: rig.digest(s) for s in before} == before


def test_company_admin_of_another_company_gets_no_foreign_board_data(rig):
    """The seed is identical in every company database, so lity's card gets a marker no other company has."""
    rig.sql("UPDATE cards SET title = 'LITY-ONLY-MARKER' WHERE id = 1", slug="lity")
    _, token = company_admin(rig, email="b@example.com", slug="salcomp")
    assert rig.call("get", "/excel/export", token, slug="lity").status_code == 403
    assert rig.call("get", "/excel/template", token, slug="lity").status_code == 403
    own = rig.call("get", "/excel/export", token, slug="salcomp")  # its own company's board, never lity's cards
    assert own.status_code == 200 and b"LITY-ONLY-MARKER" not in own.content
    ws = load_workbook(io.BytesIO(own.content)).worksheets[0]
    assert "LITY-ONLY-MARKER" not in {c for row in ws.iter_rows(values_only=True) for c in row}


def test_company_admin_unknown_board_is_404(rig):
    _, token = company_admin(rig)
    r = rig.client.get("/api/companies/lity/boards/999/excel/export", headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 404
