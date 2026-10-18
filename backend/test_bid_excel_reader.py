"""Part 22B: bid_excel_reader - every workbook is synthetic, built in memory;
no real report, no personal data, no database."""
import hashlib
import io
import os
import subprocess
import sys
import zipfile
from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook

from bid_excel_reader import (
    BID_COLUMNS,
    DEFAULT_LIMITS,
    ReaderLimits,
    Severity,
    read_bid_excel,
    read_bid_excel_file,
)

_HEADERS = [spec.header for spec in BID_COLUMNS]
# The real report's 31 columns: "Nível de Confiança" (no longer imported) sits where it always did.
HEADERS = _HEADERS[:14] + ["Nível de Confiança"] + _HEADERS[14:]


def row(**overrides):
    """A valid row (fictitious data), by header; overrides by header too."""
    values = {
        "Código": "000123-01",
        "Ano": 2026,
        "CAT": "OPEN",
        "Atualizado em": "05/03/2026",
        "Estágio": "RFQ",
        "Planta": "Planta Teste",
        "Cliente": "Cliente Ficticio",
        "Tipo de cliente": "LEAD",
        "Tipo de negócio": "CASH COW",
        "Categoria de Produto": "Charger",
        "Produto": "Produto X",
        "Solução": "ODM",
        "Nome do projeto": "Projeto Teste",
        "Descrição da prospecção": "Descricao",
        "Nível de Confiança": "LOW (10%)",
        "Volume (ano atual)": 1500,
        "Volume (próximo ano)": 3000,
        "Data de MP": "01/07/2026",
        "Revenue BRL (ano atual)": Decimal("1234.56"),
        "Revenue BRL (próximo ano)": "-",
        "Descrição do status": "Status",
        "Budget habilitado": "Não",
        "Budget": "-",
        "CPM": "Pessoa A",
        "Líder do projeto": "Pessoa B",
        "Próximo passo": "Ligar",
        "Deadline": "30/06/2026",
        "Razão": "-",
        "Status da conclusão": "-",
        "Data de criação": "01/02/2026",
        "Última edição": "-",
    }
    values.update(overrides)
    return values


def build(rows=None, headers=None, sheet_title="Leads", extra_sheets=(), formats=None) -> bytes:
    """An .xlsx with `headers` (default: the 31) in row 1 and one row per dict."""
    headers = HEADERS if headers is None else headers
    rows = [row()] if rows is None else rows
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(headers)
    for values in rows:
        ws.append([values.get(h) for h in headers] if isinstance(values, dict) else values)
    for (cell, number_format) in (formats or {}).items():
        ws[cell].number_format = number_format
    for title, sheet_rows in extra_sheets:
        extra = wb.create_sheet(title)
        for r in sheet_rows:
            extra.append(r)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def rezip(data: bytes, replace=None, add=None, drop=()) -> bytes:
    """Copy an .xlsx package, replacing/adding/removing members."""
    replace, add = replace or {}, add or {}
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            if item.filename in drop:
                continue
            content = src.read(item.filename)
            if item.filename in replace:
                content = replace[item.filename](content)
            dst.writestr(item.filename, content)
        for name, content in add.items():
            dst.writestr(name, content)
    return out.getvalue()


def with_number(text: bytes, **overrides) -> bytes:
    """A workbook whose current-year revenue cell holds exactly `text` in its XML
    (openpyxl itself writes Decimal through float, e.g. 1e+16)."""
    data = build([row(**{"Revenue BRL (ano atual)": 777, **overrides})])
    return rezip(data, replace={"xl/worksheets/sheet1.xml": lambda xml: xml.replace(b"<v>777</v>", b"<v>" + text + b"</v>")})


def inject_doctype(doctype: bytes):
    def inject(xml):
        if xml.startswith(b"<?xml"):
            head, _, rest = xml.partition(b"?>")
            return head + b"?>" + doctype + rest
        return doctype + xml
    return inject


def read(data, filename="relatorio-leads.xlsx", **limits):
    return read_bid_excel(data, filename, ReaderLimits(**limits) if limits else DEFAULT_LIMITS)


def codes(issues):
    return [issue.code for issue in issues]


def only_file_error(result):
    assert not result.ok and result.rows == () and result.sheet is None
    [error] = result.file_errors
    assert error.severity is Severity.ERROR and error.row is None
    return error.code


# --- valid file -----------------------------------------------------------------

class TestValidFile:
    def test_all_31_columns_mapped_and_typed(self):
        result = read(build([row(), row(**{"Código": "000124-01"})]))
        assert result.ok and result.file_warnings == ()
        assert result.sheet.name == "Leads"
        assert len(result.sheet.headers) == 31 and len(result.sheet.mapped_columns) == 30
        assert result.sheet.ignored_columns == ("Nível de Confiança",)  # recognised and ignored: no warning (checked above)
        assert result.summary.rows_read == 2 and result.summary.valid_rows == 2
        first = result.rows[0]
        assert first.row_number == 2 and result.rows[1].row_number == 3
        record = first.record
        assert record.code == "000123-01" and not record.code_needs_review
        assert (record.year, record.cat, record.stage) == (2026, "OPEN", "RFQ")
        assert record.customer == "Cliente Ficticio" and record.project_name == "Projeto Teste"
        assert (record.volume_current_year, record.volume_next_year) == (1500, 3000)
        assert record.revenue_current_year_cents == 123456 and record.revenue_currency == "BRL"
        assert record.next_step_due == date(2026, 6, 30)
        assert record.source_created_on == date(2026, 2, 1)
        assert record.cpm == "Pessoa A" and record.project_lead == "Pessoa B"
        assert first.raw.cells["Código"] == "000123-01"

    def test_file_metadata(self):
        data = build()
        result = read(data, filename="/some/dir/relatorio-leads.xlsx")
        assert result.file.filename == "relatorio-leads.xlsx"
        assert result.file.size_bytes == len(data)
        assert result.file.sha256 == hashlib.sha256(data).hexdigest()
        assert result.file.uncompressed_bytes > 0

    def test_cat_and_stage_kept_as_written(self):
        result = read(build([row(CAT="Open CE", **{"Estágio": "CLOSED (BY SALCOMP)"})]))
        record = result.rows[0].record
        assert (record.cat, record.stage) == ("Open CE", "CLOSED (BY SALCOMP)")

    def test_headers_match_ignoring_case_and_extra_spaces(self):
        headers = [f"  {h.upper()} " if h == "Cliente" else h for h in HEADERS]
        result = read(build([[v for v in row().values()]], headers=headers))
        assert result.ok and result.rows[0].record.customer == "Cliente Ficticio"

    def test_bid_sheet_found_among_others(self):
        result = read(build(sheet_title="Leads 2026", extra_sheets=[("Notas", [["texto livre"]])]))
        assert result.ok and result.sheet.name == "Leads 2026"
        assert result.sheet.all_sheet_names == ("Leads 2026", "Notas")

    def test_read_from_path(self, tmp_path):
        path = tmp_path / "relatorio-leads.xlsx"
        path.write_bytes(build())
        assert read_bid_excel_file(str(path)).ok
        assert only_file_error(read_bid_excel_file(str(path), ReaderLimits(max_file_bytes=10))) == "file_too_large"


# --- values ----------------------------------------------------------------------

class TestValues:
    def test_text_and_excel_dates(self):
        result = read(build([row(**{"Data de MP": datetime(2026, 7, 1), "Deadline": "28/02/2026"})]))
        record = result.rows[0].record
        assert record.production_start_date == date(2026, 7, 1)
        assert record.next_step_due == date(2026, 2, 28)

    @pytest.mark.parametrize("value", ["31/02/2026", "2026-03-05", "5/3/2026", "05/13/2026", "01/01/1850", 45000])
    def test_invalid_dates_are_row_errors(self, value):
        result = read(build([row(Deadline=value)]))
        [error] = result.rows[0].errors
        assert (error.code, error.column, error.field, error.row) == ("invalid_date", "Deadline", "next_step_due", 2)
        assert result.rows[0].record.next_step_due is None

    @pytest.mark.parametrize("value, cents", [
        (Decimal("1234.56"), 123456),
        (Decimal("0.10"), 10),
        (1234, 123400),
        ("1234.50", 123450),
    ])
    def test_revenue_exact_cents(self, value, cents):
        record = read(build([row(**{"Revenue BRL (ano atual)": value})])).rows[0].record
        assert record.revenue_current_year_cents == cents

    @pytest.mark.parametrize("text, cents", [
        (b"9999999999999999.99", 999999999999999999),   # float would give 1e+16
        (b"12345678901234.57", 1234567890123457),
        (b"0.1", 10),
        (b"1.5E3", 150000),
    ])
    def test_revenue_is_read_from_the_cell_text_never_through_float(self, text, cents):
        assert read(with_number(text)).rows[0].record.revenue_current_year_cents == cents

    def test_revenue_with_more_precision_than_cents_is_rejected(self):
        [error] = read(with_number(b"10.005")).rows[0].errors
        assert error.code == "invalid_amount" and "two decimal places" in error.message

    @pytest.mark.parametrize("value, reason", [
        (Decimal("-10"), "negative"),
        ("-0.01", "negative"),
        (Decimal("10.005"), "more than two decimal places"),
        ("1.234,56", "not a decimal number"),
        ("R$ 10", "not a decimal number"),
        (Decimal("10000000000000000"), "exceeds the maximum"),
        (datetime(2026, 1, 1), "date where an amount"),
    ])
    def test_invalid_revenue_is_a_row_error_without_the_value(self, value, reason):
        result = read(build([row(**{"Revenue BRL (próximo ano)": value})]))
        [error] = result.rows[0].errors
        assert error.code in ("invalid_amount", "invalid_type") and reason in error.message
        assert str(value) not in error.message
        assert result.rows[0].record.revenue_next_year_cents is None

    def test_dash_blank_and_empty_are_absent_not_zero(self):
        result = read(build([row(**{
            "Volume (próximo ano)": "-", "Revenue BRL (ano atual)": " - ", "Data de MP": None, "Razão": "",
        })]))
        current = result.rows[0]
        assert current.is_valid
        record = current.record
        assert record.volume_next_year is None and record.revenue_current_year_cents is None
        assert record.production_start_date is None and record.close_reason_note is None
        assert {"volume_next_year", "revenue_current_year_cents", "production_start_date", "close_reason_note"} <= set(
            current.empty_fields
        )

    def test_volume_one_is_kept_with_a_warning(self):
        result = read(build([row(**{"Volume (ano atual)": 1})]))
        current = result.rows[0]
        assert current.is_valid and current.record.volume_current_year == 1
        [warning] = current.warnings
        assert (warning.code, warning.severity, warning.field) == ("volume_one", Severity.WARNING, "volume_current_year")
        assert result.summary.rows_with_warnings == 1 and result.summary.valid_rows == 1

    @pytest.mark.parametrize("value, expected, code", [
        (Decimal("1.5E3"), 1500, None),
        ("2500", 2500, None),
        (Decimal("10.5"), None, "invalid_number"),
        (-5, None, "invalid_number"),
        ("1.500", None, "invalid_type"),
    ])
    def test_volumes_are_whole_numbers(self, value, expected, code):
        current = read(build([row(**{"Volume (próximo ano)": value})])).rows[0]
        assert current.record.volume_next_year == expected
        assert codes(current.errors) == ([code] if code else [])

    @pytest.mark.parametrize("value", ["LOW (10%)", "70%", 50, "LOW", "HIGH (150%)", Decimal("0.5")])
    def test_the_confidence_column_is_recognised_but_never_read_or_validated(self, value):
        result = read(build([row(**{"Nível de Confiança": value})]))
        assert result.rows[0].errors == () and not result.file_errors
        assert "unknown_column" not in codes(result.file_warnings)
        assert not hasattr(result.rows[0].record, "confidence_label") and not hasattr(result.rows[0].record, "confidence_percent")

    def test_year(self):
        assert read(build([row(Ano=2027)])).rows[0].record.year == 2027
        assert codes(read(build([row(Ano=1800)])).rows[0].errors) == ["invalid_number"]

    def test_numbers_in_text_columns_become_text(self):
        record = read(build([row(Produto=4521)])).rows[0].record
        assert record.product == "4521"

    def test_text_too_long(self):
        current = read(build([row(**{"Descrição da prospecção": "x" * 101})]), max_text_length=100).rows[0]
        assert codes(current.errors) == ["text_too_long"]

    def test_excel_error_value_is_a_row_error(self):
        assert codes(read(build([row(**{"Volume (ano atual)": "#N/A"})])).rows[0].errors) == ["excel_error"]

    def test_boolean_in_a_number_column_is_a_row_error(self):
        assert codes(read(build([row(**{"Volume (ano atual)": True})])).rows[0].errors) == ["invalid_type"]


class TestCodes:
    def test_leading_zeros_in_text_are_kept(self):
        record = read(build([row(**{"Código": "000007-03"})])).rows[0].record
        assert record.code == "000007-03" and not record.code_needs_review

    def test_numeric_code_with_zero_padded_format_is_rebuilt(self):
        current = read(build([row(**{"Código": 123})], formats={"A2": "000000"})).rows[0]
        assert current.record.code == "000123" and not current.record.code_needs_review
        assert "code_from_number" in codes(current.warnings)

    def test_numeric_code_without_format_is_flagged_not_padded(self):
        current = read(build([row(**{"Código": 123})])).rows[0]
        assert current.record.code == "123" and current.record.code_needs_review
        assert "code_needs_review" in codes(current.warnings) and current.is_valid

    def test_unusual_code_format_is_a_warning(self):
        current = read(build([row(**{"Código": "ABC"})])).rows[0]
        assert current.record.code == "ABC" and codes(current.warnings) == ["unexpected_code_format"]

    @pytest.mark.parametrize("field", ["Código", "CAT", "Estágio", "Cliente"])
    def test_required_values(self, field):
        current = read(build([row(**{field: "-"})])).rows[0]
        assert not current.is_valid
        [error] = current.errors
        assert (error.code, error.column) == ("required_value_missing", field)


# --- rows ------------------------------------------------------------------------

class TestRows:
    def test_an_invalid_row_does_not_stop_the_others(self):
        result = read(build([
            row(Deadline="99/99/2026"),
            row(**{"Código": "000200-01"}),
            row(**{"Revenue BRL (ano atual)": "-1", "Data de criação": "x"}),
        ]))
        assert result.ok
        assert [r.is_valid for r in result.rows] == [False, True, False]
        assert [r.row_number for r in result.rows] == [2, 3, 4]
        assert result.summary.invalid_rows == 2 and result.summary.row_errors == 3
        assert result.rows[0].record.code == "000123-01"  # the rest of an invalid row is still read

    def test_formula_without_value_is_a_row_error(self):
        # openpyxl saves formulas without a cached value, like a file never opened in Excel.
        result = read(build([row(**{"Revenue BRL (ano atual)": "=1+1"}), row(**{"Código": "000300-01"})]))
        [error] = result.rows[0].errors
        assert (error.code, error.column) == ("formula_without_value", "Revenue BRL (ano atual)")
        assert result.rows[1].is_valid

    def test_formula_with_cached_value_is_read(self):
        data = rezip(build([row(**{"Volume (ano atual)": "=10*2"})]), replace={
            "xl/worksheets/sheet1.xml": lambda xml: xml.replace(b"<f>10*2</f><v />", b"<f>10*2</f><v>20</v>"),
        })
        current = read(data).rows[0]
        assert current.is_valid and current.record.volume_current_year == 20

    def test_blank_rows_are_counted_not_returned(self):
        result = read(build([row(), [None] * 31, ["", "  "] + [None] * 29, row(**{"Código": "000400-01"})]))
        assert [r.row_number for r in result.rows] == [2, 5]
        assert result.summary.blank_rows_skipped == 2

    def test_row_of_dashes_is_reported_not_skipped(self):
        result = read(build([["-"] * 31]))
        assert result.rows[0].row_number == 2 and not result.rows[0].is_valid

    def test_values_outside_the_headers_are_reported(self):
        current = read(build([list(row().values()) + ["extra"]])).rows[0]
        assert "value_outside_headers" in codes(current.warnings)

    def test_messages_never_contain_cell_values(self):
        secret = "Cliente Confidencial SA"
        result = read(build([row(Cliente=secret, Deadline="31/02/2026", **{"Volume (ano atual)": 1})]))
        issues = result.rows[0].errors + result.rows[0].warnings
        assert issues and all(secret not in i.message and "31/02/2026" not in i.message for i in issues)


# --- structure ---------------------------------------------------------------------

class TestStructure:
    @pytest.mark.parametrize("missing", ["Código", "CAT", "Estágio", "Cliente", "Nome do projeto"])
    def test_missing_required_header_is_a_file_error(self, missing):
        headers = [h for h in HEADERS if h != missing]
        assert only_file_error(read(build([[row()[h] for h in headers]], headers=headers))) == "sheet_not_found"

    def test_missing_optional_header_is_a_warning(self):
        headers = [h for h in HEADERS if h != "Budget"]
        result = read(build([[row()[h] for h in headers]], headers=headers))
        assert result.ok and result.sheet.missing_optional_columns == ("Budget",)
        [warning] = result.file_warnings
        assert (warning.code, warning.column, warning.field) == ("missing_optional_column", "Budget", "budget")
        assert result.rows[0].record.budget is None

    def test_unknown_header_is_a_warning(self):
        headers = HEADERS + ["Coluna Nova"]
        result = read(build([list(row().values()) + ["x"]], headers=headers))
        assert result.ok and result.sheet.ignored_columns == ("Nível de Confiança", "Coluna Nova")
        assert codes(result.file_warnings) == ["unknown_column"]

    def test_duplicate_header_is_a_file_error(self):
        headers = HEADERS + ["cliente"]
        assert only_file_error(read(build([list(row().values()) + ["x"]], headers=headers))) == "duplicate_header"

    def test_blank_header_between_columns_is_a_file_error(self):
        headers = HEADERS[:5] + [None] + HEADERS[5:]
        assert only_file_error(read(build([], headers=headers))) == "blank_header"

    def test_no_bid_sheet(self):
        data = build([["a", "b"]], headers=["Outra", "Planilha"], sheet_title="Resumo")
        assert only_file_error(read(data)) == "sheet_not_found"

    def test_two_bid_sheets_are_ambiguous(self):
        data = build(extra_sheets=[("Copia", [HEADERS, list(row().values())])])
        assert only_file_error(read(data)) == "ambiguous_sheet"

    def test_headers_not_in_row_1(self):
        wb = Workbook()
        wb.active.append(["Relatorio"])
        wb.active.append(HEADERS)
        buffer = io.BytesIO()
        wb.save(buffer)
        assert only_file_error(read(buffer.getvalue())) == "sheet_not_found"

    def test_headers_without_data(self):
        assert only_file_error(read(build([]))) == "no_data_rows"

    def test_row_limit(self):
        assert read(build([row()] * 3), max_rows=3).ok
        assert only_file_error(read(build([row()] * 4), max_rows=3)) == "too_many_rows"
        assert DEFAULT_LIMITS.max_rows == 5000

    def test_column_limit(self):
        headers = HEADERS + [f"Extra {i}" for i in range(10)]
        assert only_file_error(read(build([], headers=headers), max_columns=40)) == "too_many_columns"


# --- file format and safety --------------------------------------------------------

class TestFileFormat:
    @pytest.mark.parametrize("filename, code", [
        ("relatorio.xlsm", "format_not_allowed"),
        ("relatorio.xls", "format_not_allowed"),
        ("relatorio.csv", "invalid_extension"),
        ("relatorio", "invalid_extension"),
    ])
    def test_extensions(self, filename, code):
        assert only_file_error(read(build(), filename=filename)) == code

    def test_uppercase_extension_is_accepted(self):
        assert read(build(), filename="RELATORIO.XLSX").ok

    @pytest.mark.parametrize("data, code", [
        (b"", "empty_file"),
        (b"Codigo;CAT\n1;OPEN\n", "not_xlsx"),                 # a CSV renamed to .xlsx
        (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 100, "format_not_allowed"),  # legacy .xls renamed
        (b"PK\x03\x04garbage", "corrupted_file"),
    ])
    def test_not_an_xlsx(self, data, code):
        assert only_file_error(read(data)) == code

    def test_truncated_file(self):
        data = build()
        assert only_file_error(read(data[: len(data) // 2])) == "corrupted_file"

    def test_zip_without_workbook(self):
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as z:
            z.writestr("hello.txt", "hi")
        assert only_file_error(read(out.getvalue())) == "not_xlsx"

    def test_corrupted_sheet_xml(self):
        data = rezip(build(), replace={"xl/worksheets/sheet1.xml": lambda xml: xml[: len(xml) // 2]})
        assert only_file_error(read(data)) == "corrupted_file"

    def test_macro_enabled_content_renamed_to_xlsx(self):
        macro = b"application/vnd.ms-excel.sheet.macroEnabled.main+xml"
        data = rezip(build(), replace={
            "[Content_Types].xml": lambda xml: xml.replace(
                b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml", macro
            ),
        })
        assert only_file_error(read(data)) == "format_not_allowed"

    def test_vba_project_is_refused(self):
        data = rezip(build(), add={"xl/vbaProject.bin": b"\0" * 10})
        assert only_file_error(read(data)) == "format_not_allowed"

    def test_file_size_limit(self):
        data = build()
        assert only_file_error(read(data, max_file_bytes=len(data) - 1)) == "file_too_large"

    def test_uncompressed_size_limit(self):
        assert only_file_error(read(build(), max_uncompressed_bytes=1000)) == "uncompressed_too_large"

    def test_zip_bomb_member_is_refused(self):
        data = rezip(build(), add={"xl/media/padding.bin": b"\0" * (5 * 1024 * 1024)})
        assert only_file_error(read(data)) == "suspicious_compression"

    def test_too_many_zip_entries(self):
        assert only_file_error(read(build(), max_zip_entries=3)) == "too_many_entries"

    @pytest.mark.parametrize("part", ["xl/workbook.xml", "xl/styles.xml", "xl/worksheets/sheet1.xml", "[Content_Types].xml"])
    @pytest.mark.parametrize("doctype", [
        b'<!DOCTYPE x [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;">]>',  # entity expansion
        b'<!DOCTYPE x [<!ENTITY ext SYSTEM "file:///etc/passwd">]>',                       # external entity
    ])
    def test_entity_declarations_are_refused_in_every_part(self, part, doctype):
        data = rezip(build(), replace={part: inject_doctype(doctype)})
        assert only_file_error(read(data)) == "unsafe_xml"


# --- no side effects ------------------------------------------------------------------

def test_reading_touches_no_database(tmp_path):
    """The reader imports nothing database-related, and the test database is unchanged."""
    database_path = os.environ["DATABASE_URL"].removeprefix("sqlite:///")
    before = hashlib.sha256(open(database_path, "rb").read()).hexdigest() if os.path.exists(database_path) else None
    path = tmp_path / "relatorio-leads.xlsx"
    path.write_bytes(build([row(), row(**{"Código": "000999-01"})]))
    script = (
        "import sys; sys.path.insert(0, sys.argv[1]);"
        "from bid_excel_reader import read_bid_excel_file;"
        "result = read_bid_excel_file(sys.argv[2]);"
        "assert result.ok and result.summary.rows_read == 2;"
        "loaded = [m for m in ('sqlalchemy', 'database', 'models', 'sqlite3') if m in sys.modules];"
        "print(','.join(loaded) or 'none')"
    )
    backend = os.path.dirname(os.path.abspath(__file__))
    env = {k: v for k, v in os.environ.items() if k != "DATABASE_URL"}
    out = subprocess.run(
        [sys.executable, "-I", "-c", script, backend, str(path)], capture_output=True, text=True, env=env, check=True,
    )
    assert out.stdout.strip() == "none"
    after = hashlib.sha256(open(database_path, "rb").read()).hexdigest() if os.path.exists(database_path) else None
    assert before == after
